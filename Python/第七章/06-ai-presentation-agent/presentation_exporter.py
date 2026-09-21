"""将通过领域校验的页面导出为可继续编辑的 PPTX 文件。"""

from __future__ import annotations

import os
import re
from pathlib import Path

from pptx import Presentation as PptxPresentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from presentation_aggregate import now_iso
from presentation_types import ExportRecord, OutlineVersion, PageTask, Presentation


class PresentationExporter:
    """PPTX 导出只读取当前页面产物，不会尝试补全缺失页面。"""

    def __init__(self, output_dir: Path | str | None = None) -> None:
        default_dir = Path(__file__).resolve().parent / "data" / "exports"
        self.output_dir = (
            Path(output_dir).resolve()
            if output_dir is not None
            else Path(os.getenv("EXPORT_DIR", str(default_dir))).resolve()
        )

    @staticmethod
    def _color(value: str) -> RGBColor:
        return RGBColor.from_string(value)

    @staticmethod
    def _set_text(
        shape: object,
        text: str,
        *,
        font_face: str,
        size: float,
        color: str,
        bold: bool = False,
        align: str = "left",
    ) -> None:
        text_frame = shape.text_frame  # type: ignore[attr-defined]
        text_frame.clear()
        text_frame.word_wrap = True
        paragraph = text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
        run = paragraph.runs[0]
        run.font.name = font_face
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)

    def export(
        self,
        presentation: Presentation,
        outline: OutlineVersion,
        pages: list[PageTask],
    ) -> ExportRecord:
        """生成宽屏 PPTX，并把每页 speakerNote 写进备注页。"""

        self.output_dir.mkdir(parents=True, exist_ok=True)
        pptx = PptxPresentation()
        pptx.slide_width = Inches(13.333)
        pptx.slide_height = Inches(7.5)
        pptx.core_properties.author = "DeckFlow Agent"
        pptx.core_properties.subject = presentation.requirements.topic
        pptx.core_properties.title = outline.title
        pptx.core_properties.company = "Agent Course"
        pptx.core_properties.language = "zh-CN"
        blank_layout = pptx.slide_layouts[6]

        for page in pages:
            artifact = next(
                (item for item in page.artifacts if item.artifactId == page.currentArtifactId),
                None,
            )
            if artifact is None:
                raise ValueError(f"页面 {page.pageId} 缺少当前产物。")

            is_cover = page.order == 1
            style = page.styleOverride
            layout_style = style.layoutStyle if style else presentation.theme.layoutStyle
            title_align = style.titleAlign if style else presentation.theme.titleAlign
            is_compact = (style.density if style else presentation.theme.density) == "compact"
            background = (
                style.backgroundColor
                if style
                else presentation.theme.coverBackgroundColor
                if is_cover
                else presentation.theme.backgroundColor
            )
            text_color = (
                style.textColor
                if style
                else presentation.theme.coverTextColor
                if is_cover
                else presentation.theme.textColor
            )
            muted_color = (
                style.mutedColor
                if style
                else presentation.theme.coverMutedColor
                if is_cover
                else presentation.theme.mutedColor
            )
            accent_color = style.accentColor if style else presentation.theme.accentColor
            head_font = style.headFontFace if style else presentation.theme.headFontFace
            body_font = style.bodyFontFace if style else presentation.theme.bodyFontFace

            slide = pptx.slides.add_slide(blank_layout)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = self._color(background)

            if layout_style == "top_line":
                x, y, width, height = 0, 0, 13.333, 0.12
            elif layout_style == "corner_block":
                x, y, width, height = 0.85, 0.55, 0.72, 0.1
            else:
                x, y, width, height = 0, 0, 0.18, 7.5
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
            accent.fill.solid()
            accent.fill.fore_color.rgb = self._color(accent_color)
            accent.line.color.rgb = self._color(accent_color)

            number_box = slide.shapes.add_textbox(Inches(11.7), Inches(0.45), Inches(0.8), Inches(0.4))
            self._set_text(number_box, f"{page.order:02d}", font_face="Aptos Mono", size=11, color=muted_color, align="right")

            title_box = slide.shapes.add_textbox(
                Inches(1.35 if title_align == "center" else 0.85),
                Inches(1.0 if layout_style == "corner_block" else 0.75),
                Inches(10.6),
                Inches(1.2),
            )
            self._set_text(
                title_box,
                artifact.title,
                font_face=head_font,
                size=31 if is_cover else 26,
                color=text_color,
                bold=True,
                align=title_align,
            )

            subtitle_box = slide.shapes.add_textbox(Inches(0.88), Inches(2.0), Inches(10.2), Inches(0.45))
            self._set_text(subtitle_box, artifact.subtitle, font_face=body_font, size=12, color=muted_color, align=title_align)

            bullets_box = slide.shapes.add_textbox(
                Inches(0.92),
                Inches(2.6 if is_compact else 2.75),
                Inches(10.6),
                Inches(3.4 if is_compact else 3.2),
            )
            frame = bullets_box.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            for index, item in enumerate(artifact.bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = item
                paragraph.level = 0
                paragraph.font.name = body_font
                paragraph.font.size = Pt(16 if is_compact else 18)
                paragraph.font.color.rgb = self._color(text_color)
                paragraph.space_after = Pt(10 if is_compact else 16)

            footer_box = slide.shapes.add_textbox(Inches(0.88), Inches(6.85), Inches(5.5), Inches(0.3))
            self._set_text(
                footer_box,
                f"Outline v{outline.version} · Revision {page.pageRevision}",
                font_face=body_font,
                size=9,
                color=muted_color,
            )

            # python-pptx 1.x 支持通过 notes_slide.notes_text_frame 写入讲者备注。
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes_frame.text = artifact.speakerNote

        safe_name = re.sub(r'[\\/:*?"<>|]', "-", presentation.requirements.topic)[:50]
        file_name = f"{safe_name}-v{outline.version}.pptx"
        file_path = self.output_dir / file_name
        pptx.save(file_path)
        return ExportRecord(
            fileName=file_name,
            filePath=str(file_path),
            outlineVersion=outline.version,
            pageArtifactIds=[page.currentArtifactId for page in pages if page.currentArtifactId],
            createdAt=now_iso(),
        )
