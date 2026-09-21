"""综合验证领域规则、工作流恢复、自然语言修改和本地 PPTX 导出。"""

from __future__ import annotations

from tempfile import TemporaryDirectory
import unittest

from langgraph.checkpoint.memory import MemorySaver
from pptx import Presentation as PptxPresentation

from presentation_aggregate import PresentationAggregate
from presentation_exporter import PresentationExporter
from presentation_graph import PresentationGraphService
from presentation_model import PresentationModelService
from presentation_repository import InMemoryPresentationRepository
from presentation_service import PresentationService
from presentation_types import (
    ApplyChangeInput,
    CreatePresentationInput,
    EditableRequirements,
    GeneratedPage,
    OutlineDraft,
    OutlineDraftSlide,
    ReviewOutlineInput,
)


INPUT = CreatePresentationInput(
    modelMode="replay",
    topic="企业 Agent 发布方案",
    audience="技术负责人",
    pageCount=3,
    additionalRequirements="",
    sourceText="这是一份长度足够的企业 Agent 课程资料，用于测试大纲和页面状态。",
    sourceName="source.md",
)


def create_service(output_dir: str) -> PresentationService:
    repository = InMemoryPresentationRepository()
    models = PresentationModelService()
    graph = PresentationGraphService(
        repository=repository,
        models=models,
        exporter=PresentationExporter(output_dir=output_dir),
        checkpointer=MemorySaver(),
    )
    return PresentationService(repository, graph, models)


class PresentationAggregateTests(unittest.TestCase):
    def test_only_the_current_outline_can_be_reviewed(self) -> None:
        aggregate = PresentationAggregate.create(INPUT)
        draft = OutlineDraft(
            title="企业 Agent 发布方案",
            slides=[
                OutlineDraftSlide(
                    title=f"第 {number} 页",
                    purpose=f"解释主题 {number}",
                    keyPoints=[f"重点 {number}-1", f"重点 {number}-2"],
                )
                for number in (1, 2, 3)
            ],
        )
        aggregate.add_outline(draft, None)
        aggregate.add_outline(draft.model_copy(update={"title": "企业 Agent 发布方案 v2"}), "补充风险控制")

        with self.assertRaisesRegex(ValueError, "已经过期"):
            aggregate.approve_outline(1)
        aggregate.approve_outline(2)
        self.assertEqual(aggregate.current_outline.status, "approved")

    def test_theme_changes_keep_pages_but_invalidate_previous_export(self) -> None:
        aggregate = PresentationAggregate.create(INPUT)
        draft = OutlineDraft(
            title=INPUT.topic,
            slides=[
                OutlineDraftSlide(title=f"第 {number} 页", purpose="测试", keyPoints=["重点一", "重点二"])
                for number in (1, 2, 3)
            ],
        )
        aggregate.add_outline(draft, None)
        aggregate.approve_outline(1)
        aggregate.ensure_page_tasks()
        for page in aggregate.get_current_pages():
            aggregate.start_page(page.pageId)
            aggregate.complete_page(
                page.pageId,
                GeneratedPage(title=page.title, subtitle="副标题", bullets=["要点一", "要点二"], speakerNote="备注"),
            )
        aggregate.change_theme("technology", "整体改成深色科技风")
        value = aggregate.to_model()
        self.assertEqual(value.theme.preset, "technology")
        self.assertEqual(value.theme.revision, 2)
        self.assertTrue(all(page.status == "completed" for page in value.pages))
        self.assertIsNone(value.exportRecord)


class PresentationWorkflowTests(unittest.TestCase):
    def test_replay_preserves_successful_pages_then_only_retries_the_failed_page(self) -> None:
        with TemporaryDirectory() as directory:
            service = create_service(directory)
            task = service.create(INPUT)
            self.assertEqual(task.status, "waiting_review")

            task = service.review(
                task.id,
                ReviewOutlineInput(decision="approve", outlineVersion=1, feedback=""),
            )
            self.assertEqual(task.status, "partially_completed")
            self.assertEqual([page.status for page in task.pages], ["completed", "completed", "failed"])
            self.assertEqual([page.attempts for page in task.pages], [1, 1, 1])

            task = service.continue_pages(task.id)
            self.assertEqual(task.status, "completed")
            self.assertEqual([page.attempts for page in task.pages], [1, 1, 2])
            self.assertEqual(len(task.pages[2].artifacts), 1)

    def test_page_revision_and_pptx_export(self) -> None:
        with TemporaryDirectory() as directory:
            service = create_service(directory)
            task = service.create(INPUT)
            task = service.review(task.id, ReviewOutlineInput(decision="approve", outlineVersion=1))
            task = service.continue_pages(task.id)

            task = service.revise_page(task.id, "page-2", "突出业务收益")
            page_two = next(page for page in task.pages if page.pageId == "page-2")
            self.assertEqual(page_two.pageRevision, 2)
            self.assertEqual(page_two.artifacts[-1].bullets[0], "突出业务收益")
            self.assertEqual(next(page for page in task.pages if page.pageId == "page-1").attempts, 1)

            task = service.export(task.id)
            self.assertEqual(task.status, "exported")
            self.assertIsNotNone(task.exportRecord)
            pptx = PptxPresentation(task.exportRecord.filePath)
            self.assertEqual(len(pptx.slides), 3)
            self.assertIn("本页用于", pptx.slides[0].notes_slide.notes_text_frame.text)

    def test_global_change_creates_a_new_outline_and_page_style_does_not_regenerate_content(self) -> None:
        with TemporaryDirectory() as directory:
            service = create_service(directory)
            task = service.create(INPUT)
            task = service.review(task.id, ReviewOutlineInput(decision="approve", outlineVersion=1))
            task = service.continue_pages(task.id)

            task, plan = service.apply_change(
                task.id,
                ApplyChangeInput(instruction="第 2 页使用绿色背景"),
            )
            self.assertEqual(plan.scope, "single_page_style")
            page_two = next(page for page in task.pages if page.pageId == "page-2")
            self.assertEqual(page_two.styleOverride.backgroundColor, "12372A")
            self.assertEqual(page_two.pageRevision, 1)

            task, plan = service.apply_change(
                task.id,
                ApplyChangeInput(instruction="增加一页企业 Agent 落地风险与控制方案。"),
            )
            self.assertEqual(plan.scope, "global_content")
            self.assertEqual(task.currentOutlineVersion, 2)
            self.assertEqual(task.status, "waiting_review")
            self.assertTrue(all(page.status == "stale" for page in task.pages))
            self.assertEqual(task.requirements.pageCount, 4)


if __name__ == "__main__":
    unittest.main()
