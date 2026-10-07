"""离线验证：只操作课程样本、临时文件和本地 DuckDB，不调用外部服务。"""

import io
import json
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import duckdb
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet.formula import ArrayFormula

import import_data
from table_parser import (
    COLUMNS, MAX_ROWS, build_profile, inspect_source,
    normalize_cell, normalize_rows, read_source, read_table,
)


ROOT = Path(__file__).resolve().parent
XLSX_OPTIONS = {"sheet": "销售明细", "header_row": 3}
HEADERS = [column["header"] for column in COLUMNS]
ROW = ["0001", "001001", "2026-08-03", "华东", "P01", "1", "1200", "0", "备注"]


def load(name):
    source = read_source(ROOT / "samples" / name)
    table = read_table(source, **(XLSX_OPTIONS if "workbook" in source else {}))
    normalized = normalize_rows(table)
    return source, table, normalized, build_profile(table, normalized)


def convert(name, value, extra=None, date1904=False):
    column = next(column for column in COLUMNS if column["name"] == name)
    return normalize_cell({"t": "n" if isinstance(value, (int, float)) else "s",
                           "v": value, **(extra or {})}, column, date1904)


def csv_source(row=ROW, delimiter=","):
    import csv

    text = io.StringIO(newline="")
    writer = csv.writer(text, delimiter=delimiter)
    writer.writerow(HEADERS)
    writer.writerow(row)
    return {"text": text.getvalue()}


def write_workbook(file, *, rows=None, date1904=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售明细"
    if date1904:
        workbook.epoch = CALENDAR_MAC_1904
    sheet.append(HEADERS)
    for row in rows or [ROW]:
        sheet.append(row)
    workbook.save(file)
    workbook.close()


class TableImportTests(unittest.TestCase):
    def test_excel_and_csv_values_preserve_ids_commas_and_multiline_notes(self):
        _, excel, normalized_excel, profile = load("sales-clean.xlsx")
        _, csv, normalized_csv, _ = load("sales-clean.csv")
        self.assertEqual([row["values"] for row in normalized_excel["rows"]],
                         [row["values"] for row in normalized_csv["rows"]])
        self.assertEqual(profile["status"], "ready")
        self.assertEqual(profile["rowCount"], 6)
        self.assertEqual(profile["validRowCount"], 6)
        self.assertEqual(profile["issueCount"], 0)
        self.assertEqual(len({row["values"]["order_id"] for row in normalized_excel["rows"]}), 5)
        self.assertEqual(normalized_excel["rows"][0]["values"]["line_id"], "0001")
        self.assertEqual(normalized_excel["rows"][0]["values"]["order_id"], "001001")
        self.assertEqual(normalized_csv["rows"][0]["values"]["note"], "企业客户,首次采购")
        self.assertIn("\n", normalized_csv["rows"][4]["values"]["note"])
        self.assertEqual(profile["dateRange"], {"start": "2026-08-03", "end": "2026-09-18"})
        self.assertEqual(excel["rows"][0]["sourceRow"], 4)
        self.assertEqual(csv["rows"][0]["sourceRow"], 2)

    def test_workbook_selection_header_and_preview(self):
        source = read_source(ROOT / "samples/sales-clean.xlsx")
        preview = inspect_source(source)
        self.assertEqual([item["sheet"] for item in preview], ["使用说明", "销售明细", "商品信息"])
        self.assertEqual([item["range"] for item in preview], ["A1:A9", "A1:I9", "A1:C3"])
        self.assertEqual(preview[1]["preview"][3][2], "2026-08-03")
        self.assertEqual(preview[1]["preview"][3][6], "2,400.50")
        with self.assertRaisesRegex(ValueError, "指定有效工作表"):
            read_table(source)
        with self.assertRaisesRegex(ValueError, "表头不匹配"):
            read_table(source, sheet="商品信息")
        with self.assertRaisesRegex(ValueError, "合并单元格"):
            read_table(source, sheet="销售明细")
        for header_row in (0, -1, 1.5, True, "3"):
            with self.subTest(header_row=header_row), self.assertRaisesRegex(ValueError, "正整数"):
                read_table(source, sheet="销售明细", header_row=header_row)

    def test_duplicate_missing_and_reordered_headers(self):
        source = read_source(ROOT / "samples/sales-clean.xlsx")
        sheet = source["workbook"]["Sheets"]["销售明细"]
        sheet["A3"]["v"] = "订单编号"
        with self.assertRaisesRegex(ValueError, "表头不匹配"):
            read_table(source, **XLSX_OPTIONS)
        sheet["A3"]["v"] = "不支持的字段"
        with self.assertRaisesRegex(ValueError, "表头不匹配"):
            read_table(source, **XLSX_OPTIONS)
        sheet["A3"]["v"] = "明细编号"
        sheet["!merges"].append({"s": {"r": 3, "c": 0}, "e": {"r": 4, "c": 0}})
        with self.assertRaisesRegex(ValueError, "合并单元格"):
            read_table(source, **XLSX_OPTIONS)
        source = csv_source()
        source["text"] = ",".join(reversed(HEADERS)) + "\n" + ",".join(reversed(ROW))
        self.assertEqual(normalize_rows(read_table(source))["rows"][0]["values"]["order_id"], "001001")

    def test_issue_rows_and_cached_formula_are_preserved_not_used(self):
        _, table, normalized, profile = load("sales-issues.xlsx")
        self.assertEqual(profile["rowCount"], 10)
        self.assertEqual(profile["status"], "needs_review")
        self.assertEqual(profile["validRowCount"], 5)
        self.assertEqual(profile["issueCount"], 5)
        self.assertEqual([issue["sourceRow"] for issue in normalized["issues"]], [11, 12, 13, 4, 10])
        self.assertEqual(table["rows"][-1]["cells"]["paid_amount"]["f"], "G4+G5")
        self.assertEqual(table["rows"][-1]["cells"]["paid_amount"]["v"], 2799.5)
        self.assertIsNone(normalized["rows"][-1]["values"]["paid_amount"])
        self.assertIsNone(normalized["rows"][7]["values"]["sold_at"])
        self.assertIsNone(normalized["rows"][8]["values"]["paid_amount"])
        self.assertEqual(normalized["issues"][2]["code"], "FORMULA_REQUIRES_REVIEW")
        self.assertEqual(profile["dateRange"]["end"], "2026-09-25")

    def test_money_is_decimal_text_without_silent_correction(self):
        for value, expected in (("2,400.50", "2400.50"), (0, "0.00"), ("0.1", "0.10"),
                                ("9999999999999999.99", "9999999999999999.99"), (" 12 ", "12.00")):
            with self.subTest(value=value):
                self.assertEqual(convert("paid_amount", value), expected)
        for value in (None, "", "2,40", "2万元", "-1", "1.234", "10000000000000000", True):
            with self.subTest(value=value), self.assertRaises(ValueError):
                convert("paid_amount", value)
        self.assertIsNone(convert("note", ""))
        for extra in ({"f": "SUM(A1:A2)"}, {"F": "G4:G5"}):
            with self.assertRaisesRegex(ValueError, "公式"):
                convert("paid_amount", 12, extra)
        with self.assertRaisesRegex(ValueError, "错误值"):
            convert("paid_amount", 7, {"t": "e"})

    def test_dates_real_calendar_and_both_excel_epochs(self):
        for value, expected, epoch in (("2024-02-29", "2024-02-29", False),
                                      (1, "1900-01-01", False), (59, "1900-02-28", False),
                                      (61, "1900-03-01", False), (0, "1904-01-01", True)):
            with self.subTest(value=value, epoch=epoch):
                self.assertEqual(convert("sold_at", value, date1904=epoch), expected)
        for value in ("2026-09-31", "2026-02-29", "09/10/2026", 0, 60, 1.5, -1, 2_958_466):
            with self.subTest(value=value), self.assertRaises(ValueError):
                convert("sold_at", value)

    def test_raw_date_serial_is_not_lost_by_openpyxl_conversion(self):
        with tempfile.TemporaryDirectory() as directory:
            file = Path(directory) / "epoch.xlsx"
            row = list(ROW)
            row[2] = datetime(2026, 8, 3)
            write_workbook(file, rows=[row], date1904=True)
            source = read_source(file)
            table = read_table(source, sheet="销售明细")
            self.assertTrue(table["date1904"])
            self.assertIsInstance(table["rows"][0]["cells"]["sold_at"]["v"], int)
            self.assertEqual(normalize_rows(table)["rows"][0]["values"]["sold_at"], "2026-08-03")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "销售明细"
            sheet.append(HEADERS)
            sheet.append([*ROW[:2], 60, *ROW[3:]])
            sheet["C2"].number_format = "yyyy-mm-dd"
            workbook.save(file)
            workbook.close()
            table = read_table(read_source(file), sheet="销售明细")
            self.assertEqual(table["rows"][0]["cells"]["sold_at"]["v"], 60)
            self.assertIsNone(normalize_rows(table)["rows"][0]["values"]["sold_at"])

    def test_formula_without_cache_and_array_result_region_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            file = Path(directory) / "formula.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "销售明细"
            sheet.append(HEADERS)
            sheet.append(ROW)
            sheet.append(["0002", *ROW[1:]])
            sheet["G2"] = ArrayFormula(ref="G2:G3", text="=SUM(F2:F3)")
            sheet["G3"] = None
            workbook.save(file)
            workbook.close()
            table = read_table(read_source(file), sheet="销售明细")
            self.assertEqual(table["rows"][1]["cells"]["paid_amount"]["F"], "G2:G3")
            normalized = normalize_rows(table)
            self.assertEqual([issue["code"] for issue in normalized["issues"]],
                             ["FORMULA_REQUIRES_REVIEW", "FORMULA_REQUIRES_REVIEW"])

    def test_id_and_quantity_rules(self):
        self.assertEqual(convert("order_id", "001001"), "001001")
        with self.assertRaisesRegex(ValueError, "文本"):
            convert("order_id", 1001)
        self.assertEqual(convert("quantity", "2"), 2)
        self.assertEqual(convert("quantity", 2.0), 2)
        self.assertEqual(convert("quantity", 1_000_000), 1_000_000)
        for value in (0, -1, 1.5, "两个", "02", 1_000_001, True):
            with self.subTest(value=value), self.assertRaises(ValueError):
                convert("quantity", value)

    def test_excessive_refund_keeps_values_and_reports_business_issue(self):
        _, table, _, _ = load("sales-clean.xlsx")
        table["rows"][0]["cells"]["refund_amount"]["v"] = "2400.51"
        normalized = normalize_rows(table)
        self.assertEqual(normalized["rows"][0]["values"]["refund_amount"], "2400.51")
        self.assertFalse(normalized["rows"][0]["isValid"])
        self.assertEqual(normalized["issues"][0]["code"], "REFUND_EXCEEDS_PAYMENT")

    def test_csv_delimiters_blank_records_and_wrong_column_counts(self):
        for delimiter in (";", "\t"):
            with self.subTest(delimiter=delimiter):
                source = csv_source(delimiter=delimiter)
                self.assertEqual(len(read_table(source, delimiter=delimiter)["rows"]), 1)
                with self.assertRaisesRegex(ValueError, "表头不匹配"):
                    read_table(source)
        source = read_source(ROOT / "samples/sales-clean.csv")
        self.assertEqual(read_table({**source, "text": source["text"] + "\r\n"})["skippedBlankRows"], 1)
        with self.assertRaisesRegex(ValueError, "列数"):
            read_table({**source, "text": source["text"] + "多余,列\r\n"})
        with self.assertRaisesRegex(ValueError, "没有工作表"):
            read_table(source, sheet="销售明细")
        with self.assertRaisesRegex(ValueError, "分隔符"):
            read_table(source, delimiter="|")
        with self.assertRaisesRegex(ValueError, "CSV 格式"):
            read_table({"text": ",".join(HEADERS) + '\n"unterminated'})

    def test_empty_sheet_and_header_without_data(self):
        with tempfile.TemporaryDirectory() as directory:
            file = Path(directory) / "empty.xlsx"
            workbook = Workbook()
            workbook.active.title = "空表"
            workbook.save(file)
            workbook.close()
            self.assertEqual(inspect_source(read_source(file))[0],
                             {"sheet": "空表", "range": "空表", "preview": []})
            with self.assertRaisesRegex(ValueError, "工作表为空"):
                read_table(read_source(file), sheet="空表")
        with self.assertRaisesRegex(ValueError, "没有数据"):
            read_table({"text": ",".join(HEADERS) + "\n"})

    def test_csv_escaped_quotes_work_but_quotes_inside_unquoted_fields_fail(self):
        row = [*ROW[:-1], '客户说："请保留原文"']
        table = read_table(csv_source(row))
        self.assertEqual(table["rows"][0]["cells"]["note"]["v"], row[-1])
        source = {"text": ",".join(HEADERS) + "\n" + ",".join(row)}
        with self.assertRaisesRegex(ValueError, "CSV 格式"):
            read_table(source)

    def test_csv_preview_does_not_use_default_small_field_limit_or_read_past_four_records(self):
        source = {"text": "x" * 140_000 + "\nsecond\nthird\nfourth\n\"unterminated"}
        preview = inspect_source(source)[0]["preview"]
        self.assertEqual(len(preview), 4)
        self.assertEqual(len(preview[0][0]), 140_000)

    def test_source_rejects_empty_wrong_encoding_old_format_and_size(self):
        with tempfile.TemporaryDirectory() as directory:
            for name, data, pattern in (
                ("empty.csv", b"", "非空"), ("bad.csv", b"\xff\xfe\xff", "UTF-8"),
                ("old.xls", b"old", "只支持"), ("fake.xlsx", b"not an xlsx", "XLSX"),
                ("large.csv", b"\x00" * (5 * 1024 * 1024 + 1), "5 MiB"),
                ("broken.xlsx", b"PKnot-a-zip", "XLSX"),
            ):
                file = Path(directory) / name
                file.write_bytes(data)
                with self.subTest(name=name), self.assertRaisesRegex(ValueError, pattern):
                    read_source(file)
            with self.assertRaisesRegex(ValueError, "非空"):
                read_source(directory)
            with self.assertRaises(FileNotFoundError):
                read_source(Path(directory) / "absent.csv")

    def test_row_column_and_record_limits(self):
        source = csv_source()
        row = ",".join(ROW) + "\n"
        source["text"] = ",".join(HEADERS) + "\n" + row * MAX_ROWS
        self.assertEqual(len(read_table(source)["rows"]), MAX_ROWS)
        with self.assertRaisesRegex(ValueError, "10000"):
            read_table({"text": source["text"] + row})
        long_row = [*ROW[:-1], "a" * 100_002]
        with self.assertRaisesRegex(ValueError, "100000"):
            read_table(csv_source(long_row))
        source = read_source(ROOT / "samples/sales-clean.xlsx")
        source["workbook"]["Sheets"]["销售明细"]["!ref"] = "A1:AY9"
        with self.assertRaisesRegex(ValueError, "50 列"):
            read_table(source, **XLSX_OPTIONS)
        source["workbook"]["Sheets"]["销售明细"]["!ref"] = "A1:I10004"
        with self.assertRaisesRegex(ValueError, "10000"):
            read_table(source, **XLSX_OPTIONS)

    def test_profile_without_legal_dates_and_optional_nulls(self):
        _, table, _, _ = load("sales-clean.xlsx")
        for row in table["rows"]:
            row["cells"]["sold_at"]["v"] = "2026-09-31"
        normalized = normalize_rows(table)
        profile = build_profile(table, normalized)
        self.assertIsNone(profile["dateRange"])
        self.assertEqual(profile["status"], "needs_review")
        self.assertEqual(profile["validRowCount"], 0)
        self.assertEqual(len(profile["samples"]), 3)
        notes = next(column for column in profile["columns"] if column["name"] == "note")
        self.assertEqual(notes["nullCount"], 3)

    def test_source_and_version_hashes_match_node_contract(self):
        expected = {
            "sales-clean.xlsx": ("b29d79f345812d399a04373e109f90f5861f4121929975465e5516b8312b6b4e",
                                 "457f9fcb6e2fc41a327eae9c22a636b28f1d8edb3df129a61a3ead4b60086494"),
            "sales-clean.csv": ("061c97571b5e548a63d060576267fa1febd2bdbd81ae652387ba4e9b02370227",
                                "8f9aee49de20d952083e32175251cd8e161a0014d8c02db28b1cfc8fff7d9406"),
            "sales-issues.xlsx": ("9c9bdd4f7f49948be6d01ebfe40fc438357595cbef6ba735df449ba536effdc5",
                                  "3cf95833470e74d1d6589b2a37339d5eb0864c78a3735913c8cd1370957b1f7e"),
        }
        with tempfile.TemporaryDirectory() as directory:
            for name, (source_hash, version) in expected.items():
                source, table, normalized, profile = load(name)
                saved = import_data.save_dataset(source, table, normalized, profile, output_root=directory)
                self.assertEqual(source["sha256"], source_hash)
                self.assertEqual(saved["version"], version)

    def test_saved_database_reopens_with_exact_amounts_and_raw_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            source, table, normalized, profile = load("sales-clean.xlsx")
            saved = import_data.save_dataset(source, table, normalized, profile, output_root=directory)
            output = Path(saved["directory"])
            self.assertEqual({file.name for file in output.iterdir()},
                             {"source.xlsx", "data.duckdb", "profile.json", "issues.json"})
            self.assertEqual((output / "source.xlsx").read_bytes(), source["bytes"])
            with duckdb.connect(str(output / "data.duckdb"), read_only=True) as database:
                self.assertEqual(database.execute("""SELECT count(*), count(DISTINCT order_id),
                    sum(paid_amount)::VARCHAR, sum(refund_amount)::VARCHAR FROM sales""").fetchone(),
                    (6, 5, "6396.50", "299.00"))
                raw = json.loads(database.execute(
                    "SELECT cells_json FROM sales_raw WHERE source_row = 4").fetchone()[0])
                self.assertEqual(raw, table["rows"][0]["cells"])
                self.assertEqual(database.execute("SELECT count(*) FROM import_issues").fetchone()[0], 0)
            repeated = import_data.save_dataset(source, table, normalized, profile, output_root=directory)
            self.assertTrue(repeated["reused"])
            self.assertEqual(repeated["directory"], saved["directory"])
            metadata = json.loads((output / "profile.json").read_text(encoding="utf-8"))
            self.assertNotIn("directory", metadata)
            self.assertNotIn("reused", metadata)
            self.assertEqual(metadata["samples"], normalized["rows"][:3])

    def test_new_version_keeps_issues_and_does_not_overwrite_clean_data(self):
        with tempfile.TemporaryDirectory() as directory:
            clean = load("sales-clean.xlsx")
            dirty = load("sales-issues.xlsx")
            first = import_data.save_dataset(*clean, output_root=directory)
            second = import_data.save_dataset(*dirty, output_root=directory)
            self.assertNotEqual(first["version"], second["version"])
            self.assertEqual(json.loads((Path(first["directory"]) / "profile.json").read_text())["status"], "ready")
            with duckdb.connect(str(Path(second["directory"]) / "data.duckdb"), read_only=True) as database:
                self.assertEqual(database.execute(
                    "SELECT count(*), count(*) FILTER (WHERE NOT is_valid) FROM sales").fetchone(), (10, 5))
                self.assertEqual(database.execute("SELECT count(*) FROM sales_raw").fetchone()[0], 10)
                self.assertEqual(database.execute("SELECT count(*) FROM import_issues").fetchone()[0], 5)
                raw = json.loads(database.execute(
                    "SELECT cells_json FROM sales_raw WHERE source_row = 13").fetchone()[0])
                self.assertEqual(raw["paid_amount"]["v"], 2799.5)
                self.assertEqual(raw["paid_amount"]["f"], "G4+G5")
            for dataset_id in ("../escape", "A", "a" * 41, "含中文", "a/b"):
                with self.subTest(dataset_id=dataset_id), self.assertRaisesRegex(ValueError, "datasetId"):
                    import_data.save_dataset(*clean, output_root=directory, dataset_id=dataset_id)

    def test_changed_import_rule_changes_version_even_when_values_match(self):
        with tempfile.TemporaryDirectory() as directory:
            source, table, normalized, profile = load("sales-clean.xlsx")
            first = import_data.save_dataset(source, table, normalized, profile, output_root=directory)
            second = import_data.save_dataset(source, {**table, "delimiter": ";"}, normalized, profile,
                                              output_root=directory)
            self.assertNotEqual(first["version"], second["version"])
            self.assertEqual(first["sourceHash"], second["sourceHash"])

    def test_database_failure_rolls_back_all_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            _, table, normalized, _ = load("sales-clean.xlsx")
            normalized = deepcopy(normalized)
            normalized["rows"][1]["values"]["paid_amount"] = "invalid decimal"
            file = Path(directory) / "broken.duckdb"
            with self.assertRaises(duckdb.Error):
                import_data.save_database(file, table, normalized)
            with duckdb.connect(str(file), read_only=True) as database:
                self.assertEqual(database.execute("SELECT count(*) FROM information_schema.tables").fetchone()[0], 0)

    def test_publish_failure_removes_temporary_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            args = load("sales-clean.xlsx")
            with patch.object(import_data, "save_database", side_effect=RuntimeError("test failure")):
                with self.assertRaisesRegex(RuntimeError, "test failure"):
                    import_data.save_dataset(*args, output_root=directory)
            self.assertEqual(list((Path(directory) / "sales-demo").iterdir()), [])

    def test_directory_publish_failure_also_removes_database_and_source_copy(self):
        with tempfile.TemporaryDirectory() as directory:
            args = load("sales-clean.xlsx")
            with patch.object(Path, "rename", side_effect=OSError("test rename failure")):
                with self.assertRaisesRegex(OSError, "test rename failure"):
                    import_data.save_dataset(*args, output_root=directory)
            self.assertEqual(list((Path(directory) / "sales-demo").iterdir()), [])

    def test_existing_corrupt_metadata_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            args = load("sales-clean.xlsx")
            saved = import_data.save_dataset(*args, output_root=directory)
            file = Path(saved["directory"]) / "profile.json"
            file.write_text('{"version":"wrong"}', encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "元数据不一致"):
                import_data.save_dataset(*args, output_root=directory)
            self.assertEqual(file.read_text(), '{"version":"wrong"}')
            self.assertFalse(any(path.name.startswith(".import-") for path in file.parent.parent.iterdir()))

    def test_cli_inspect_is_read_only_and_import_works_from_other_cwd(self):
        with tempfile.TemporaryDirectory() as directory:
            stdout = io.StringIO()
            with patch.object(import_data, "PROJECT_DIR", Path(directory)), redirect_stdout(stdout):
                import_data.main(["inspect", str(ROOT / "samples/sales-clean.xlsx")])
            self.assertIn("销售明细", stdout.getvalue())
            self.assertEqual(list(Path(directory).iterdir()), [])
            script = "import import_data; from pathlib import Path; " + (
                f"import_data.PROJECT_DIR=Path({directory!r}); "
                f"import_data.main(['import',{str(ROOT / 'samples/sales-clean.csv')!r}])"
            )
            # 显式 Python 目录只供导入当前小节；运行时不读取另一语言目录或环境文件。
            environment = {"PYTHONPATH": str(ROOT), "PYTHONDONTWRITEBYTECODE": "1"}
            process = subprocess.run([sys.executable, "-c", script], cwd=directory, env=environment,
                                     capture_output=True, text=True, timeout=20)
            self.assertEqual(process.returncode, 0, process.stderr)
            self.assertIn("状态：ready", process.stdout)
            self.assertIn("明细行数：6", process.stdout)
            outputs = list((Path(directory) / "outputs/sales-demo").iterdir())
            self.assertEqual(len(outputs), 1)


if __name__ == "__main__":
    unittest.main()
