"""解析销售明细，保留原始位置，再按已确认的字段规则整理与检查。"""

import csv
import hashlib
import io
import math
import posixpath
import re
import stat
from collections import Counter
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel


MAX_ROWS = 10_000
PARSER_VERSION = "sales-import-v1"
COLUMNS = [
    {"name": "line_id", "header": "明细编号", "type": "VARCHAR", "kind": "id",
     "description": "一行一笔订单商品明细，明细编号唯一"},
    {"name": "order_id", "header": "订单编号", "type": "VARCHAR", "kind": "id",
     "description": "同一订单可对应多条明细，保留前导零"},
    {"name": "sold_at", "header": "销售日期", "type": "DATE", "kind": "date",
     "description": "销售发生日期"},
    {"name": "region", "header": "区域", "type": "VARCHAR", "kind": "text",
     "description": "销售所属区域"},
    {"name": "product_id", "header": "商品编号", "type": "VARCHAR", "kind": "id",
     "description": "商品编号，对应原文件中的商品信息表"},
    {"name": "quantity", "header": "数量", "type": "INTEGER", "kind": "integer",
     "description": "本条明细的商品数量"},
    {"name": "paid_amount", "header": "实付金额（元）", "type": "DECIMAL(18,2)", "kind": "money",
     "description": "人民币元，本条明细实付金额，未扣退款"},
    {"name": "refund_amount", "header": "退款金额（元）", "type": "DECIMAL(18,2)", "kind": "money",
     "description": "人民币元，单列退款金额，0 表示无退款"},
    {"name": "note", "header": "备注", "type": "VARCHAR", "kind": "text", "optional": True,
     "description": "原始备注，不作为程序指令执行"},
]
XML_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def hash_value(value):
    return hashlib.sha256(value.encode("utf-8") if isinstance(value, str) else value).hexdigest()


def _text(value):
    """保留源示例的文本转换语义，避免数值 2.0 被误判为小数数量。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and math.isfinite(value):
        if value == 0:
            return "0"
        if 1e-6 <= abs(value) < 1e21:
            text = format(Decimal(str(value)), "f")
            return text.rstrip("0").rstrip(".") if "." in text else text
        return re.sub(r"e([+-])0+(\d+)$", r"e\1\2", str(value))
    return str(value)


def _blank(value):
    return value is None or isinstance(value, str) and value.strip() == ""


def calendar_date(year, month, day):
    """校验真实日历日期，不把 9 月 31 日自动滚到 10 月。"""
    if not 1900 <= year <= 9999:
        raise ValueError("日期不存在或超出本例范围。")
    try:
        return date(year, month, day).isoformat()
    except ValueError as error:
        raise ValueError("日期不存在或超出本例范围。") from error


def _excel_date(serial, date1904):
    if not 0 <= serial <= 2_958_465:
        raise ValueError("Excel 日期序列值无效。")
    # 1900 系统的 0 对应“1 月 0 日”，60 对应不存在的“2 月 29 日”。
    # 不采用 openpyxl 将 59/60 映射到同一天的结果，避免非法日期被悄悄修正。
    if not date1904 and serial in (0, 60):
        raise ValueError("日期不存在或超出本例范围。")
    try:
        day = date(1904, 1, 1) + timedelta(days=serial) if date1904 else (
            date(1899, 12, 31) + timedelta(days=serial - (1 if serial > 60 else 0))
        )
    except (ValueError, OverflowError) as error:
        raise ValueError("日期不存在或超出本例范围。") from error
    return calendar_date(day.year, day.month, day.day)


def _display_cell(cell, date1904):
    """预览使用课程样本的显示格式；类型转换始终读取原值 v，不读取显示文本 w。"""
    value, kind, number_format = cell["v"], cell["t"], cell["z"]
    if kind == "b":
        return "TRUE" if value else "FALSE"
    if kind != "n":
        return _text(value)
    if number_format.lower() == "yyyy-mm-dd" and isinstance(value, (int, float)):
        try:
            return _excel_date(int(value), date1904)
        except ValueError:
            return _text(value)
    if number_format in ("#,##0.00", "0.00"):
        return format(Decimal(str(value)), ",.2f" if "," in number_format else ".2f")
    return _text(value)


def _read_workbook(data):
    """用 openpyxl 读取工作簿，并补回原始日期序列值与公式缓存供追溯。"""
    workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    date1904 = workbook.epoch == CALENDAR_MAC_1904
    result = {"SheetNames": workbook.sheetnames, "Sheets": {},
              "Workbook": {"WBProps": {"date1904": date1904}}}
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
            relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {item.get("Id"): item.get("Target") for item in relationships}
            for item in workbook_xml.findall("s:sheets/s:sheet", XML_NS):
                name = item.get("name")
                target = targets[item.get(f"{{{REL_NS}}}id")]
                member = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
                    posixpath.join("xl", target)
                )
                worksheet_xml = ET.fromstring(archive.read(member))
                worksheet = workbook[name]
                cells = {}
                array_ranges = []
                # openpyxl 会自动把日期数值变成 datetime；直接读取 XML 的 v 保留原序列值。
                # 公式的 v 是上次保存的缓存，只作原始资料保存，后面仍会拒绝参与业务计算。
                for node in worksheet_xml.findall("s:sheetData/s:row/s:c", XML_NS):
                    address = node.get("r")
                    original = worksheet[address]
                    raw_type = node.get("t", "n")
                    raw_value = node.findtext("s:v", default=None, namespaces=XML_NS)
                    formula = node.find("s:f", XML_NS)
                    if raw_type in ("s", "inlineStr", "str"):
                        value = original.value if raw_type != "str" else raw_value
                        kind, value = "s", value if value is not None else ""
                    elif raw_type == "b":
                        kind, value = "b", raw_value == "1"
                    elif raw_type == "e":
                        kind, value = "e", raw_value
                    elif raw_type == "d":
                        kind, value = "n", to_excel(original.value, workbook.epoch)
                    elif raw_value not in (None, ""):
                        kind, value = "n", float(raw_value)
                        if value.is_integer():
                            value = int(value)
                    elif formula is not None:
                        kind, value = "n", None
                    else:
                        # 只有样式、没有值的单元格对应空单元格，不伪造 0 或格式化文本。
                        continue
                    cell = {"t": kind, "v": value, "z": original.number_format}
                    if value is not None:
                        cell["w"] = _display_cell(cell, date1904)
                    if formula is not None:
                        if isinstance(original.value, str) and original.value.startswith("="):
                            cell["f"] = original.value[1:]
                        elif formula.text:
                            cell["f"] = formula.text
                        if formula.get("t") == "array" and formula.get("ref"):
                            array_ranges.append(formula.get("ref"))
                    cells[address] = cell
                # 数组公式不仅标记左上角，还要覆盖结果区域；无缓存值也不能当作普通空白跳过。
                for reference in array_ranges:
                    left, top, right, bottom = range_boundaries(reference)
                    for row in range(top, bottom + 1):
                        for column in range(left, right + 1):
                            address = f"{get_column_letter(column)}{row}"
                            cells.setdefault(address, {"t": "z", "v": None})["F"] = reference
                merges = [{"s": {"r": merge.min_row - 1, "c": merge.min_col - 1},
                           "e": {"r": merge.max_row - 1, "c": merge.max_col - 1}}
                          for merge in worksheet.merged_cells.ranges]
                sheet = {**cells, "!merges": merges}
                if cells:
                    # 仅有样式的空单元格不扩大有效范围；与源示例的 !ref 一致。
                    bounds = [range_boundaries(address) for address in cells]
                    left, top = min(b[0] for b in bounds), min(b[1] for b in bounds)
                    right, bottom = max(b[2] for b in bounds), max(b[3] for b in bounds)
                    sheet["!ref"] = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
                result["Sheets"][name] = sheet
    finally:
        workbook.close()
    return result


def read_source(file_path):
    """读取受限大小的本地样本；CSV 只接收 UTF-8，避免乱码后继续导入。"""
    file_path = Path(file_path)
    info = file_path.stat()
    if not stat.S_ISREG(info.st_mode) or not 0 < info.st_size <= 5 * 1024 * 1024:
        raise ValueError("请选择非空且不超过 5 MiB 的本地文件。")
    extension = file_path.suffix.lower()
    if extension not in (".xlsx", ".csv"):
        raise ValueError("本例只支持 .xlsx 和 UTF-8 CSV。")
    data = file_path.read_bytes()
    source = {"filePath": str(file_path), "bytes": data, "extension": extension, "sha256": hash_value(data)}
    if extension == ".xlsx":
        if data[:2] != b"PK":
            raise ValueError("文件不是有效的 XLSX 压缩包。")
        try:
            source["workbook"] = _read_workbook(data)
        except (BadZipFile, KeyError, ET.ParseError) as error:
            raise ValueError("文件不是有效的 XLSX 压缩包。") from error
    else:
        try:
            # UTF-8 BOM 不属于第一列表头；保留原文件字节用于 Hash 和独立副本。
            source["text"] = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("CSV 编码不是 UTF-8，请从 Excel/WPS 重新导出为 CSV UTF-8。") from error
    return source


def _csv_records(text, delimiter=",", limit=None, max_record_size=None):
    # newline="" 保留引号内换行，strict=True 拒绝不完整引号，不使用 splitlines() 拆记录。
    text = text.removeprefix("\ufeff")
    stream = io.StringIO(text, newline="")
    reader = csv.reader(stream, delimiter=delimiter, strict=True)
    # csv 默认的单字段上限小于本例文件上限；正式导入的记录上限仍在下面单独检查。
    previous_limit = csv.field_size_limit()
    csv.field_size_limit(5 * 1024 * 1024)
    records = []
    position = 0
    try:
        for row in reader:
            # Python strict 模式仍允许未加引号字段中出现双引号；原案例拒绝该格式。
            # 只检查本次已读记录，inspect 不会提前解析第四条之后的内容。
            raw_record = text[position:stream.tell()]
            position = stream.tell()
            quoted, field_start, index = False, True, 0
            while index < len(raw_record):
                char = raw_record[index]
                if quoted:
                    if char == '"':
                        if index + 1 < len(raw_record) and raw_record[index + 1] == '"':
                            index += 1
                        else:
                            quoted = False
                elif char == '"':
                    if not field_start:
                        raise ValueError("CSV 格式无效：双引号必须从字段开头开始。")
                    quoted = True
                    field_start = False
                elif char == delimiter:
                    field_start = True
                else:
                    field_start = False
                index += 1
            # csv.reader 把完全空行返回为 []；原示例将它保留为一列空字符串。
            row = row or [""]
            if max_record_size is not None and sum(len(value) for value in row) > max_record_size:
                raise ValueError("CSV 单条记录超过 100000 字符。")
            records.append(row)
            if limit is not None and len(records) >= limit:
                break
    except csv.Error as error:
        raise ValueError(f"CSV 格式无效：{error}") from error
    finally:
        csv.field_size_limit(previous_limit)
    return records


def inspect_source(source):
    """预览工作表或 CSV 前四条记录，先确定目标表与表头；不执行导入。"""
    if "workbook" not in source:
        return [{"sheet": "CSV", "preview": _csv_records(source["text"], limit=4)}]
    result = []
    for name in source["workbook"]["SheetNames"]:
        sheet = source["workbook"]["Sheets"][name]
        reference = sheet.get("!ref")
        preview = []
        if reference:
            _, _, right, bottom = range_boundaries(reference)
            for row in range(1, min(bottom, 4) + 1):
                values = [sheet.get(f"{get_column_letter(column)}{row}", {}).get("w")
                          for column in range(1, min(right, 50) + 1)]
                while values and values[-1] is None:
                    values.pop()
                preview.append(values)
        result.append({"sheet": name, "range": reference or "空表", "preview": preview})
    return result


def read_table(source, *, sheet=None, header_row=1, delimiter=","):
    """按指定工作表和表头读取，统一保留单元格地址、类型、原值、格式和公式。"""
    if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
        raise ValueError("headerRow 必须是从 1 开始的正整数。")
    date1904 = False
    if "workbook" in source:
        workbook = source["workbook"]
        if not sheet or sheet not in workbook["Sheets"]:
            raise ValueError(f"请指定有效工作表：{'、'.join(workbook['SheetNames'])}。")
        worksheet = workbook["Sheets"][sheet]
        if not worksheet.get("!ref"):
            raise ValueError("选中的工作表为空。")
        _, _, right, bottom = range_boundaries(worksheet["!ref"])
        if bottom > MAX_ROWS + header_row or right > 50:
            raise ValueError("本例最多读取 10000 条数据、50 列。")
        # 允许表头上方的标题合并，拒绝表头和数据区域的合并，避免字段错位。
        if any(merge["e"]["r"] >= header_row - 1 for merge in worksheet.get("!merges", [])):
            raise ValueError("表头或数据区域包含合并单元格，请先整理成单行表头、每行一条明细。")
        date1904 = bool(workbook.get("Workbook", {}).get("WBProps", {}).get("date1904"))
        matrix = []
        for row in range(1, bottom + 1):
            cells = []
            for column in range(1, right + 1):
                address = f"{get_column_letter(column)}{row}"
                original = worksheet.get(address, {})
                cell = {"address": address, "t": original.get("t", "z"), "v": original.get("v")}
                cell.update({key: original[key] for key in ("w", "z", "f", "F") if key in original})
                cells.append(cell)
            matrix.append(cells)
    else:
        if sheet:
            raise ValueError("CSV 没有工作表，不需要 --sheet。")
        if delimiter not in (",", ";", "\t"):
            raise ValueError("分隔符只支持逗号、分号或 tab。")
        records = _csv_records(source["text"], delimiter, max_record_size=100_000)
        if len(records) > MAX_ROWS + header_row:
            raise ValueError("本例最多读取 10000 条 CSV 数据。")
        matrix = [[{"address": f"记录{row + 1}/列{column + 1}", "t": "s", "v": value}
                   for column, value in enumerate(record)] for row, record in enumerate(records)]

    header = matrix[header_row - 1] if header_row <= len(matrix) else []
    names = [_text(cell.get("v")).strip() for cell in header]
    if (len(names) != len(COLUMNS) or len(set(names)) != len(names)
            or any(column["header"] not in names for column in COLUMNS)):
        expected = "、".join(column["header"] for column in COLUMNS)
        raise ValueError(f"表头不匹配。当前：{'、'.join(names)}。需要：{expected}。请检查工作表、表头行或分隔符。")

    rows, skipped = [], 0
    for source_row, cells in enumerate(matrix[header_row:], start=header_row + 1):
        # 只跳过完全空白且没有公式的记录，跳过数量进入数据概览。
        if all(_blank(cell.get("v")) and not cell.get("f") and not cell.get("F") for cell in cells):
            skipped += 1
            continue
        if len(cells) != len(names):
            raise ValueError(f"第 {source_row} 条记录的列数与表头不同。")
        rows.append({"sourceRow": source_row,
                     "cells": {column["name"]: cells[names.index(column["header"])] for column in COLUMNS}})
    if not rows:
        raise ValueError("表头后面没有数据。")
    return {"sheet": sheet, "headerRow": header_row, "delimiter": delimiter, "date1904": date1904,
            "rows": rows, "skippedBlankRows": skipped,
            "positionKind": "Excel 行号" if "workbook" in source else "CSV 记录序号（包含表头）"}


def normalize_cell(cell, column, date1904=False):
    """按已确认的字段规则转换；不猜测单位、不修补缺失值、不使用公式缓存。"""
    if cell.get("f") or cell.get("F"):
        raise ValueError("该单元格含公式。请核实并导出数值版本；本例不采用可能过期的公式缓存。")
    if cell.get("t") == "e":
        raise ValueError("该单元格包含 Excel 错误值。")
    value = cell.get("v")
    if _blank(value):
        if column.get("optional"):
            return None
        raise ValueError("必需字段为空。")
    text = _text(value).strip()
    kind = column["kind"]
    if kind == "id":
        if not isinstance(value, str):
            raise ValueError("编号必须以文本保存，避免前导零或长编号丢失。")
        return text
    if kind == "text":
        return text
    if kind == "date":
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if not math.isfinite(value) or int(value) != value:
                raise ValueError("销售日期只接收整日，不自动舍弃时间。")
            return _excel_date(int(value), date1904)
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", text):
            raise ValueError("日期请使用 YYYY-MM-DD，避免日/月顺序歧义。")
        return calendar_date(*(int(part) for part in text.split("-")))
    if kind == "integer":
        if not re.fullmatch(r"[1-9][0-9]*", text) or int(text) > 1_000_000:
            raise ValueError("数量必须是 1 到 1000000 的整数。")
        return int(text)
    if not re.fullmatch(r"(?:0|[1-9][0-9]*|[1-9][0-9]{0,2}(?:,[0-9]{3})+)(?:\.[0-9]{1,2})?", text):
        raise ValueError("金额须为非负数、最多两位小数；支持千分位，不自动换算元和万元。")
    integer, _, fraction = text.replace(",", "").partition(".")
    if len(integer) > 16:
        raise ValueError("金额超出 DECIMAL(18,2) 范围。")
    return f"{integer}.{fraction.ljust(2, '0')}"


def normalize_rows(table):
    """收集全部问题，保留问题行；重复明细的每一行都标记，交由人工核对。"""
    rows, issues = [], []

    def add_issue(row, field, code, message):
        issues.append({"sourceRow": row["sourceRow"], "field": field, "code": code, "message": message})

    for row in table["rows"]:
        values = {}
        for column in COLUMNS:
            cell = row["cells"][column["name"]]
            try:
                values[column["name"]] = normalize_cell(cell, column, table["date1904"])
            except ValueError as error:
                # 字段失败不丢弃整行；置空并保存原始行号、字段和具体原因。
                values[column["name"]] = None
                code = "FORMULA_REQUIRES_REVIEW" if cell.get("f") or cell.get("F") else "INVALID_VALUE"
                add_issue(row, column["name"], code, str(error))
        paid, refund = values["paid_amount"], values["refund_amount"]
        # 十进制文本转成整数“分”比较，避免二进制浮点误差。
        if paid is not None and refund is not None and int(refund.replace(".", "")) > int(paid.replace(".", "")):
            add_issue(row, "refund_amount", "REFUND_EXCEEDS_PAYMENT", "退款金额超过本条明细实付金额，请核对。")
        rows.append({"sourceRow": row["sourceRow"], "values": values})

    counts = Counter(row["values"]["line_id"] for row in rows if row["values"]["line_id"])
    for row in rows:
        line_id = row["values"]["line_id"]
        if counts[line_id] > 1:
            add_issue(row, "line_id", "DUPLICATE_LINE_ID",
                      f"明细编号 {line_id} 重复。请核实，程序没有删除记录。")
    invalid_rows = {issue["sourceRow"] for issue in issues}
    return {"rows": [{**row, "isValid": row["sourceRow"] not in invalid_rows} for row in rows], "issues": issues}


def build_profile(table, normalized):
    """汇总字段、行数、合法日期与前三条样例；完整明细继续留在本地。"""
    # 合法日期来自能转换的日期字段；不等于只统计 isValid 为 true 的整行。
    dates = sorted(row["values"]["sold_at"] for row in normalized["rows"] if row["values"]["sold_at"])
    return {
        "status": "needs_review" if normalized["issues"] else "ready",
        "rowCount": len(normalized["rows"]),
        "validRowCount": sum(row["isValid"] for row in normalized["rows"]),
        "issueCount": len(normalized["issues"]),
        "skippedBlankRows": table["skippedBlankRows"],
        "dateRange": {"start": dates[0], "end": dates[-1]} if dates else None,
        "grain": "一行一笔订单商品明细；line_id 应唯一，order_id 可以重复。",
        "amountPolicy": "人民币元，paid_amount 未扣退款；refund_amount 单列，空值不等于 0。",
        "columns": [{**column, "nullCount": sum(row["values"][column["name"]] is None
                                                for row in normalized["rows"])} for column in COLUMNS],
        "samples": normalized["rows"][:3],
    }
